"""Round 18, task 3b: two real defects behind a score that looked like an artifact.

`testPPT_various2.pptx` scored 0.9767 with `NotesMasterHeader` and
`NotesMasterFooter` missing. Round 17 blamed the oracle for welding the two
words together. Half right, and the wrong half mattered: the oracle *does*
separate paragraphs, and the store really did contain the welded token
`NotesMasterHeaderNotesMasterFooter`. The weld was the reader's, in
`core/forensics/office_package.py`, which joined every ``a:t`` in a notes part
with no separator at all.

The second defect was found while measuring the first: a notes page was
attributed to the slide whose *number matched its part name*. In this deck
`notesSlide2.xml` holds "Notes for slide3", because the page carries the
number it was created with. The relationship part is the authority.

Fixtures are generated with python-pptx (the reader's own dependency) so the
packages are real ones, then the notes part is rewritten to the shape each test
needs. The corpus file is read, never written.
"""
import os
import shutil
import zipfile

import pytest

from core.forensics.office_package import PackageLimits, extract_package_features

CORPUS = os.path.join(os.path.dirname(__file__), "..", "..", "test-doc")
DECK = os.path.join(CORPUS, "Presentations", "testPPT_various2.pptx")

A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'


# --------------------------------------------------------------------------- #
# Fixture helpers
# --------------------------------------------------------------------------- #
def _make_deck(path, notes_paragraphs=("NotesMasterHeader", "NotesMasterFooter")):
    """A one-slide deck whose notes page has the given paragraphs."""
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Only slide"
    slide.notes_slide.notes_text_frame.text = "\n".join(notes_paragraphs)
    prs.save(str(path))
    return path


def _rewrite_parts(path, replace=None, rename=None, drop=()):
    """Rewrite selected parts of an OOXML package, in place."""
    replace = replace or {}
    rename = rename or {}
    with zipfile.ZipFile(path) as zf:
        items = [(name, zf.read(name)) for name in zf.namelist()]
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, payload in items:
            if name in drop:
                continue
            new_name = rename.get(name, name)
            zf.writestr(new_name, replace.get(name, payload))
    tmp.replace(path)
    return path


def _notes_part_names(path):
    with zipfile.ZipFile(path) as zf:
        return [n for n in zf.namelist()
                if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml")]


def _notes_xml(*paragraphs):
    body = "".join(
        "<a:p>" + "".join(f"<a:r><a:t>{run}</a:t></a:r>" for run in runs) + "</a:p>"
        for runs in paragraphs)
    return ('<?xml version="1.0" encoding="UTF-8"?>'
            f'<p:notes xmlns:p="http://schemas.openxmlformats.org/presentationml/'
            f'2006/main" xmlns:a="{A_NS}"><p:cSld><p:spTree><p:sp><p:nvSpPr>'
            '<p:cNvPr id="2" name="Notes Placeholder"/>'
            '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr/></p:nvSpPr>'
            '<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/>'
            f'{body}</p:txBody></p:sp></p:spTree></p:cSld></p:notes>')


def _read_deck(path):
    from reader_file.readers.read_office import OfficeFileReader

    result = OfficeFileReader().read_file({"path": str(path), "extension": ".pptx"})
    return result.get("content", result)


# --------------------------------------------------------------------------- #
# 1. Two placeholders must never become one word
# --------------------------------------------------------------------------- #
def test_adjacent_notes_placeholders_are_separated(tmp_path):
    deck = _make_deck(tmp_path / "separated.pptx")
    features = extract_package_features(deck, limits=PackageLimits())
    texts = [entry["text"] for entry in features["presentation"]["notes_text"]]
    assert texts
    joined = "\n".join(texts)
    assert "NotesMasterHeader" in joined and "NotesMasterFooter" in joined
    assert "NotesMasterHeaderNotesMasterFooter" not in joined, joined


def test_runs_in_one_paragraph_still_join_without_a_space(tmp_path):
    """A placeholder split across runs is one word, not two."""
    deck = _make_deck(tmp_path / "runs.pptx")
    part = _notes_part_names(deck)[0]
    _rewrite_parts(deck, replace={part: _notes_xml(("NotesMaster", "Header"))})
    features = extract_package_features(deck, limits=PackageLimits())
    assert features["presentation"]["notes_text"][0]["text"] == "NotesMasterHeader"


def test_each_paragraph_is_its_own_line(tmp_path):
    deck = _make_deck(tmp_path / "paras.pptx")
    part = _notes_part_names(deck)[0]
    _rewrite_parts(deck, replace={
        part: _notes_xml(("NotesMasterHeader",), ("NotesMasterFooter",))})
    features = extract_package_features(deck, limits=PackageLimits())
    assert features["presentation"]["notes_text"][0]["text"] == (
        "NotesMasterHeader\nNotesMasterFooter")


def test_the_corpus_deck_has_no_welded_note_tokens():
    if not os.path.exists(DECK):
        pytest.skip("corpus file missing")
    features = extract_package_features(DECK, limits=PackageLimits())
    joined = "\n".join(entry["text"]
                       for entry in features["presentation"]["notes_text"])
    assert "NotesMasterHeaderNotesMasterFooter" not in joined


# --------------------------------------------------------------------------- #
# 2. A notes page belongs to the slide its relationship names
# --------------------------------------------------------------------------- #
def _mismatch_deck(tmp_path):
    """A deck whose notes page is named notesSlide2 but belongs to slide 1."""
    deck = _make_deck(tmp_path / "mismatch.pptx", notes_paragraphs=("Notes for slide three",))
    part = _notes_part_names(deck)[0]                    # ppt/notesSlides/notesSlide1.xml
    rels = "ppt/notesSlides/_rels/notesSlide1.xml.rels"
    _rewrite_parts(deck, rename={part: "ppt/notesSlides/notesSlide2.xml",
                                 rels: "ppt/notesSlides/_rels/notesSlide2.xml.rels"})
    return deck


def test_notes_go_to_the_slide_the_relationship_names(tmp_path):
    deck = _mismatch_deck(tmp_path)
    content = _read_deck(deck)
    assert "slides" in content, content
    texts = [element.get("text") or ""
             for slide in content["slides"]
             for element in slide.get("elements") or []]
    texts += [str(slide.get("notes") or "") for slide in content["slides"]]
    assert any("Notes for slide three" in text for text in texts), texts


def test_the_corpus_deck_puts_its_notes_page_on_the_right_slide():
    if not os.path.exists(DECK):
        pytest.skip("corpus file missing")
    content = _read_deck(DECK)
    notes = {int(entry["slide_number"]): entry["notes"]
             for entry in content.get("presentation_notes") or []}
    assert "Notes for slide3" in notes.get(3, ""), notes
    assert "Notes for slide3" not in notes.get(2, ""), notes


def test_notes_text_is_in_the_forensic_channel_when_no_slide_matches(tmp_path):
    """Text from a part with no slide of that number is still not dropped."""
    deck = _make_deck(tmp_path / "orphan.pptx", notes_paragraphs=("orphan note",))
    part = _notes_part_names(deck)[0]
    # A notes master is read by the same walk but belongs to no slide.
    _rewrite_parts(deck, rename={part: part.replace("notesSlides/notesSlide1",
                                                    "notesMasters/notesMaster2")})
    content = _read_deck(deck)
    blob = (content.get("forensic_text") or "") + str(content.get("xml_backfill"))
    assert "orphan note" in blob, blob[:400]
