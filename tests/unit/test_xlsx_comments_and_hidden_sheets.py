"""Spreadsheets: the parts of a workbook a value-only read never sees.

Two categories of content sat outside the cell grid and were therefore invisible
to extraction *and* to search:

* **cell comments** - a reviewer's note, a correction, the reason a figure is
  what it is. openpyxl hangs them off cells (`cell.comment`), and
  `iter_rows(values_only=True)` cannot see them.
* **hidden sheets** - `sheet_state` is `hidden` or `veryHidden`. The content is
  part of the workbook and is frequently the working the author did not want on
  screen. The sheet is read like any other; the state is recorded beside the
  content so a reader of the result knows what was on display.

The comment goes into the stored text through the forensic channel, formatted
with its address, so a search for the note's wording finds the artifact.
"""

import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

openpyxl = pytest.importorskip("openpyxl", reason="openpyxl not installed")

from reader_file.readers.read_office import OfficeFileReader  # noqa: E402


def _content(result):
    if isinstance(result.get("content"), dict) and "elements" not in result:
        return result["content"]
    return result


@pytest.fixture()
def commented_workbook(tmp_path):
    workbook = openpyxl.Workbook()
    visible = workbook.active
    visible.title = "Summary"
    visible["A1"] = "Revenue"
    visible["B1"] = 1200

    hidden = workbook.create_sheet("Working")
    hidden["A1"] = "adjustment carried forward"
    hidden["A2"] = 40
    hidden.sheet_state = "hidden"

    deep = workbook.create_sheet("Scratch")
    deep["A1"] = "very hidden working note"
    deep.sheet_state = "veryHidden"

    from openpyxl.comments import Comment

    visible["B1"].comment = Comment("Third-party figure, see invoice 44", "A. Examiner")
    path = tmp_path / "book.xlsx"
    workbook.save(str(path))
    return path


def test_cell_comment_is_stored_with_its_address_and_author(commented_workbook):
    content = _content(OfficeFileReader().read_file(
        {"path": str(commented_workbook), "extension": ".xlsx"}))
    comments = content["sheets"]["Summary"]["comments"]
    assert comments == [{"cell": "B1", "author": "A. Examiner",
                         "text": "Third-party figure, see invoice 44"}]
    assert content["comment_count"] == 1

    # and it reaches the text the pipeline stores
    from pipeline.storage_pipeline import StoragePipeline

    text = object.__new__(StoragePipeline)._extract_text_from_content(content)
    assert "invoice 44" in text
    assert "Summary!B1" in text


def test_hidden_and_very_hidden_sheets_are_read_and_reported(commented_workbook):
    content = _content(OfficeFileReader().read_file(
        {"path": str(commented_workbook), "extension": ".xlsx"}))
    assert content["sheets"]["Working"]["data"][0][0] == "adjustment carried forward"
    assert content["sheets"]["Scratch"]["data"][0][0] == "very hidden working note"
    assert content["hidden_sheets"] == ["Working", "Scratch"]
    assert content["sheets"]["Summary"]["hidden"] is False


def test_a_workbook_without_comments_stores_none(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "just data"
    path = tmp_path / "plain.xlsx"
    workbook.save(str(path))

    content = _content(OfficeFileReader().read_file(
        {"path": str(path), "extension": ".xlsx"}))
    assert "comment_count" not in content
    assert not content.get("hidden_sheets")
    assert "comments" not in content["sheets"]["Sheet"]
