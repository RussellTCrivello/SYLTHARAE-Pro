"""The committed category workbook is held to its own rules.

``classification/domain_data.xlsx`` is the source of truth for taxonomy
content: worksheets are categories, column ``Term`` is one word, column
``Keyword`` is a phrase. This module re-reads the real file on every test run,
so an edit that breaks a rule fails here rather than at import time.

The counts are asserted as exact numbers on purpose. They change only when
somebody edits the workbook on purpose, and then the changed number is the
review: a silent drop of a category or a thousand keywords cannot slip through
as "the tests still pass".
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from apps.importing.validators import workbook_rules as rules  # noqa: E402

WORKBOOK = ROOT / "classification" / "domain_data.xlsx"

#: Measured after the two-column conversion; see
#: ``audit_report/r22_workbook/REPORT.md`` for the command that produced them.
EXPECTED_CATEGORIES = 56
EXPECTED_TERMS = 3096
EXPECTED_DISTINCT_TERMS = 2882
EXPECTED_KEYWORDS = 7958
EXPECTED_DISTINCT_KEYWORDS = 7625


pytestmark = pytest.mark.skipif(not WORKBOOK.exists(),
                                reason="classification/domain_data.xlsx not present")


@pytest.fixture(scope="module")
def workbook():
    return rules.read(str(WORKBOOK))


def test_the_workbook_is_present_and_two_column(workbook):
    assert workbook.layout == "two-column", (
        "the committed workbook must carry Term/Keyword headers; a legacy "
        "single-column file loses the distinction between the two"
    )


def test_the_workbook_has_no_rule_violations(workbook):
    assert workbook.errors == [], "\n".join(
        f"{v.sheet}!{v.cell}: [{v.code}] {v.message}" for v in workbook.errors)


def test_the_counts_match_the_reviewed_baseline(workbook):
    assert len(workbook.categories) == EXPECTED_CATEGORIES
    assert workbook.total_terms() == EXPECTED_TERMS
    assert workbook.distinct_terms() == EXPECTED_DISTINCT_TERMS
    assert workbook.total_keywords() == EXPECTED_KEYWORDS
    assert workbook.distinct_keywords() == EXPECTED_DISTINCT_KEYWORDS


def test_every_term_is_one_word(workbook):
    for category in workbook.categories:
        for term in category.term_texts:
            assert len(term.split()) == 1, f"{category.name}: {term!r}"


def test_every_keyword_is_at_least_two_words(workbook):
    for category in workbook.categories:
        for keyword in category.keyword_texts:
            assert len(keyword.split()) >= 2, f"{category.name}: {keyword!r}"


def test_the_helper_sheets_sit_outside_the_category_sheets(workbook):
    """The reader takes every non-underscore sheet as a category.

    ``Master Index`` therefore has to be recognised by name, and the help
    sheets must not be mistaken for categories. This pins the arrangement so a
    reordering of the file cannot quietly turn a helper sheet into a category.
    """
    book = openpyxl.load_workbook(WORKBOOK, read_only=True)
    names = book.sheetnames
    assert names[0] == rules.README_SHEET
    assert names[-2:] == [rules.OWNERS_SHEET, rules.INDEX_SHEET]
    category_sheets = [n for n in names if n not in (
        rules.README_SHEET, rules.OWNERS_SHEET, rules.INDEX_SHEET)]
    assert category_sheets == [category.name for category in workbook.categories]


def test_the_category_order_is_the_documented_one(workbook):
    """Category order is domain order, and it is stable.

    The order reaches the UI (the browse tree lists categories as stored), so
    a reordering of the sheets is a visible change and is asserted as such
    rather than left to review.
    """
    assert [category.name for category in workbook.categories][:5] == [
        "Networks", "Cybersecurity", "Telecommunications",
        "AI & Machine Learning", "Software Development"]


def test_the_owner_sheet_only_names_real_categories(workbook):
    known = {category.name.casefold() for category in workbook.categories}
    unknown = {owner for owner in workbook.owners.values()
               if owner.casefold() not in known}
    assert unknown == set()


def test_the_readme_and_index_sheets_exist(workbook):
    book = openpyxl.load_workbook(WORKBOOK, read_only=True)
    assert rules.README_SHEET in book.sheetnames
    assert rules.INDEX_SHEET in book.sheetnames


def test_the_index_sheet_lists_every_category_and_term(workbook):
    """``Master Index`` is derived; it must agree with the category sheets."""
    book = openpyxl.load_workbook(WORKBOOK, read_only=True)
    rows = [row for row in book[rules.INDEX_SHEET].iter_rows(min_row=2, values_only=True)
            if row and row[0] and row[2]]
    assert [row[1] for row in rows][0] in ("term", "keyword")
    indexed_terms = {(str(r[0]).strip().casefold(), str(r[2]).strip().casefold())
                     for r in rows if str(r[1]).strip().lower() == "term"}
    indexed_keywords = {(str(r[0]).strip().casefold(), str(r[2]).strip().casefold())
                        for r in rows if str(r[1]).strip().lower() == "keyword"}
    expected_terms = {(c.name.casefold(), t.casefold())
                      for c in workbook.categories for t in c.term_texts}
    expected_keywords = {(c.name.casefold(), k.casefold())
                         for c in workbook.categories for k in c.keyword_texts}
    assert expected_terms - indexed_terms == set(), "terms missing from Master Index"
    assert indexed_terms - expected_terms == set(), "Master Index rows with no term"
    assert expected_keywords - indexed_keywords == set(), "keywords missing from Master Index"
    assert indexed_keywords - expected_keywords == set(), "Master Index rows with no keyword"


def test_the_index_word_count_agrees_with_the_rules(workbook):
    """The index stores a word count; it must equal what the rule counts."""
    book = openpyxl.load_workbook(WORKBOOK, read_only=True)
    mismatches = []
    for row in book[rules.INDEX_SHEET].iter_rows(min_row=2, values_only=True):
        if not row or not row[2]:
            continue
        entry, stated = str(row[2]), row[3]
        counted = len(entry.split())
        if stated != counted:
            mismatches.append((entry, stated, counted))
    assert mismatches == [], f"word counts disagree: {mismatches[:5]}"


def test_no_sheet_has_a_stray_third_column(workbook):
    """Content outside Term/Keyword would be invisible to the importer."""
    book = openpyxl.load_workbook(WORKBOOK, read_only=True)
    extras = {}
    for name in book.sheetnames:
        if name in (rules.README_SHEET, rules.INDEX_SHEET, rules.OWNERS_SHEET):
            continue
        sheet = book[name]
        for row in sheet.iter_rows(min_row=1, min_col=3, values_only=True):
            for value in row:
                if value not in (None, ""):
                    extras.setdefault(name, []).append(value)
                    break
    assert extras == {}, f"unexpected columns: {extras}"


def test_the_committed_workbook_round_trips(tmp_path, workbook):
    """A write of the committed content must reproduce it rule-for-rule."""
    target = tmp_path / "roundtrip.xlsx"
    rules.write(str(target), workbook)
    again = rules.read(str(target))
    assert again.errors == []
    assert [c.name for c in again.categories] == [c.name for c in workbook.categories]
    for before, after in zip(workbook.categories, again.categories):
        assert after.term_texts == before.term_texts, before.name
        assert after.keyword_texts == before.keyword_texts, before.name
