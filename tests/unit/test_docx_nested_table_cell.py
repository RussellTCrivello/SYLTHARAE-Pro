"""A DOCX table cell that contains a nested table must not come back empty.

python-docx's ``cell.text`` walks the cell's own ``w:p`` children only. In a
Word document it is perfectly legal for a cell's entire content to be another
table (a form layout, a nested grid), and for nested tables to repeat. Word
shows that text; ``cell.text`` returns ``""`` for it, and the reader was storing
``cell.text`` - so the text vanished between the file and storage.

Measured on the round-15 corpus over 698 files:
``test-doc/Documents/testInstrLink.docx`` is a body that is one table nested to
depth 4; the reader stored **0 of 22** document tokens. ``017097.docx`` (depth
5) stored 8%. Neither is an oracle artefact - ``python-docx`` itself reports the
cell as empty, which is exactly the reading the reader inherited.
"""

import sys
from pathlib import Path

import pytest

# python-docx's namespace map has no "mc", "wps" or "v" prefix, so the text-box
# fixtures spell the declarations out.
NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
)

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

pytest.importorskip("docx", reason="python-docx not installed")


def _read(tmp_path, build):
    from docx import Document

    from reader_file.readers.read_office import OfficeFileReader

    doc = Document()
    build(doc)
    path = tmp_path / "nested.docx"
    doc.save(str(path))

    result = OfficeFileReader().read_file(
        {"path": str(path), "effective_extension": ".docx"})
    assert not result.get("error"), result.get("error")
    return result


def _flatten(result):
    """Every string the reader produced, wherever it is in the result."""
    out = []

    def walk(node):
        if isinstance(node, str):
            out.append(node)
        elif isinstance(node, dict):
            for value in node.values():
                walk(value)
        elif isinstance(node, (list, tuple)):
            for value in node:
                walk(value)

    walk(result.get("elements") or [])
    return "\n".join(out)


def test_nested_table_cell_text_survives(tmp_path):
    def build(doc):
        outer = doc.add_table(rows=1, cols=1)
        inner = outer.cell(0, 0).add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "InnerLeft"
        inner.cell(0, 1).text = "InnerRight"

    text = _flatten(_read(tmp_path, build))
    assert "InnerLeft" in text, text
    assert "InnerRight" in text, text


def test_nested_table_text_keeps_document_order(tmp_path):
    """The outer cell's own paragraph comes before the nested table."""
    def build(doc):
        outer = doc.add_table(rows=1, cols=1)
        cell = outer.cell(0, 0)
        cell.paragraphs[0].text = "BeforeNested"
        inner = cell.add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "InsideNested"
        cell.add_paragraph("AfterNested")

    text = _flatten(_read(tmp_path, build))
    assert text.index("BeforeNested") < text.index("InsideNested"), text
    assert text.index("InsideNested") < text.index("AfterNested"), text


def test_two_levels_of_nesting(tmp_path):
    def build(doc):
        outer = doc.add_table(rows=1, cols=1)
        middle = outer.cell(0, 0).add_table(rows=1, cols=1)
        inner = middle.cell(0, 0).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "DeepText"

    text = _flatten(_read(tmp_path, build))
    assert "DeepText" in text, text


def test_flat_table_rows_unchanged(tmp_path):
    """The fix must not change the shape of an ordinary table."""
    def build(doc):
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A1"
        table.cell(0, 1).text = "B1"
        table.cell(1, 0).text = "A2"
        table.cell(1, 1).text = "B2"

    result = _read(tmp_path, build)
    tables = [el for el in result["elements"] if el.get("type") == "table"]
    assert len(tables) == 1, tables
    assert tables[0]["rows"] == [["A1", "B1"], ["A2", "B2"]], tables[0]["rows"]


# --------------------------------------------------------------------------- #
# Text boxes: content that is not in the paragraph's own run stream
# --------------------------------------------------------------------------- #
def _texts(result):
    return [el.get("text") or "" for el in result["elements"]]


def test_text_box_text_is_stored(tmp_path):
    """A caption inside a DrawingML text box is not part of paragraph.text."""
    from docx.oxml import parse_xml

    def build(doc):
        paragraph = doc.add_paragraph("BodyParagraph")
        paragraph._p.append(parse_xml(
            '<w:r %s><w:drawing><wp:inline><a:graphic><a:graphicData>'
            '<wps:wsp><wps:txbx><w:txbxContent>'
            '<w:p><w:r><w:t>BoxedCaption</w:t></w:r></w:p>'
            '</w:txbxContent></wps:txbx></wps:wsp>'
            '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'
            % NS))

    texts = _texts(_read(tmp_path, build))
    assert "BodyParagraph" in texts, texts
    assert "BoxedCaption" in texts, texts


def test_text_box_text_follows_its_paragraph(tmp_path):
    from docx.oxml import parse_xml

    def build(doc):
        doc.add_paragraph("First")
        paragraph = doc.add_paragraph("Second")
        paragraph._p.append(parse_xml(
            '<w:r %s><w:drawing><wp:inline><a:graphic><a:graphicData>'
            '<wps:wsp><wps:txbx><w:txbxContent>'
            '<w:p><w:r><w:t>InsideBox</w:t></w:r></w:p>'
            '</w:txbxContent></wps:txbx></wps:wsp>'
            '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'
            % NS))
        doc.add_paragraph("Third")

    texts = _texts(_read(tmp_path, build))
    assert texts.index("Second") < texts.index("InsideBox") < texts.index("Third"), texts


def test_alternate_content_fallback_is_not_duplicated(tmp_path):
    """mc:Choice and mc:Fallback hold the same box; store it once."""
    from docx.oxml import parse_xml

    def build(doc):
        paragraph = doc.add_paragraph("Anchored")
        paragraph._p.append(parse_xml(
            '<w:r %s><mc:AlternateContent>'
            '<mc:Choice Requires="wps"><w:drawing><wp:inline><a:graphic>'
            '<a:graphicData><wps:wsp><wps:txbx><w:txbxContent>'
            '<w:p><w:r><w:t>OnlyOnce</w:t></w:r></w:p>'
            '</w:txbxContent></wps:txbx></wps:wsp></a:graphicData>'
            '</a:graphic></wp:inline></w:drawing></mc:Choice>'
            '<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>'
            '<w:p><w:r><w:t>OnlyOnce</w:t></w:r></w:p>'
            '</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>'
            '</mc:AlternateContent></w:r>'
            % NS))

    texts = _texts(_read(tmp_path, build))
    assert texts.count("OnlyOnce") == 1, texts
